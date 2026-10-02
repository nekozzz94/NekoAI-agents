#!/usr/bin/env python3
"""
Parse infrastructure requirements from various sources:
  - Jira ticket (via API)
  - Markdown file (.md)
  - DOCX file (.docx) — requires python-docx
  - Existing Terraform code directory

Outputs a structured JSON object that the Terraform Provisioner agent uses.

Usage:
  python3 parse_requirements.py --jira INFRA-123
  python3 parse_requirements.py --md /path/to/design.md
  python3 parse_requirements.py --docx /path/to/design.docx
  python3 parse_requirements.py --tf-dir ~/terraform-repos/platform-gcp-terraform/stacks/prod
"""

import argparse
import json
import os
import re
import subprocess
import sys
from pathlib import Path


GCP_RESOURCE_KEYWORDS = {
    "google_compute_instance":        ["compute instance", "vm", "virtual machine", "gce", "compute engine"],
    "google_sql_database_instance":   ["cloud sql", "sql database", "postgres", "mysql", "database instance"],
    "google_storage_bucket":          ["gcs", "storage bucket", "cloud storage", "bucket"],
    "google_compute_network":         ["vpc", "network", "virtual network"],
    "google_compute_subnetwork":      ["subnet", "subnetwork"],
    "google_compute_firewall":        ["firewall", "firewall rule", "ingress", "egress"],
    "google_pubsub_topic":            ["pubsub", "pub/sub", "topic", "message queue"],
    "google_pubsub_subscription":     ["subscription", "pubsub subscription"],
    "google_redis_instance":          ["redis", "memorystore", "cache"],
    "google_service_account":         ["service account", "sa", "iam sa"],
    "google_container_cluster":       ["gke", "kubernetes", "k8s", "container cluster"],
    "google_dns_managed_zone":        ["dns", "dns zone", "managed zone"],
    "google_compute_router_nat":      ["cloud nat", "nat gateway"],
    "google_compute_network_peering": ["vpc peering", "network peering"],
}

ENV_KEYWORDS = {
    "dev": ["dev", "development", "develop"],
    "staging": ["staging", "stage", "uat", "qa"],
    "prod": ["prod", "production", "live"],
}


def detect_resource_types(text: str) -> list[str]:
    text_lower = text.lower()
    found = []
    for resource_type, keywords in GCP_RESOURCE_KEYWORDS.items():
        if any(kw in text_lower for kw in keywords):
            found.append(resource_type)
    return found


def detect_environment(text: str) -> str | None:
    text_lower = text.lower()
    for env, keywords in ENV_KEYWORDS.items():
        if any(kw in text_lower for kw in keywords):
            return env
    return None


def parse_markdown(path: str) -> dict:
    text = Path(path).read_text()
    resource_types = detect_resource_types(text)
    env = detect_environment(text)

    # Extract title from first H1 or H2
    title_match = re.search(r"^#{1,2}\s+(.+)$", text, re.MULTILINE)
    title = title_match.group(1).strip() if title_match else Path(path).stem

    # Extract Jira key if mentioned
    jira_key = None
    jira_match = re.search(r"\b([A-Z]+-\d+)\b", text)
    if jira_match:
        jira_key = jira_match.group(1)

    return {
        "source": "markdown",
        "source_path": str(path),
        "title": title,
        "jira_key": jira_key,
        "environment": env,
        "resource_types": resource_types,
        "raw_text": text[:2000],
    }


def parse_docx(path: str) -> dict:
    try:
        from docx import Document
    except ImportError:
        print("python-docx not installed. Run: pip install python-docx", file=sys.stderr)
        sys.exit(1)

    doc = Document(path)
    text = "\n".join(p.text for p in doc.paragraphs)
    resource_types = detect_resource_types(text)
    env = detect_environment(text)

    title = doc.paragraphs[0].text.strip() if doc.paragraphs else Path(path).stem
    jira_match = re.search(r"\b([A-Z]+-\d+)\b", text)
    jira_key = jira_match.group(1) if jira_match else None

    return {
        "source": "docx",
        "source_path": str(path),
        "title": title,
        "jira_key": jira_key,
        "environment": env,
        "resource_types": resource_types,
        "raw_text": text[:2000],
    }


def parse_jira(ticket_key: str) -> dict:
    base_url = os.environ.get("JIRA_BASE_URL", "")
    if not base_url:
        print("Set JIRA_BASE_URL environment variable", file=sys.stderr)
        sys.exit(1)

    try:
        token = subprocess.check_output(
            ["security", "find-generic-password", "-s", "JiraAPIToken", "-w"],
            stderr=subprocess.DEVNULL,
        ).decode().strip()
    except subprocess.CalledProcessError:
        print("Keychain item 'JiraAPIToken' not found. Add it with:\n"
              "  security add-generic-password -s JiraAPIToken -a '' -w '<token>'",
              file=sys.stderr)
        sys.exit(1)

    import urllib.request

    url = f"{base_url}/rest/api/3/issue/{ticket_key}?fields=summary,description,labels,priority,status"
    req = urllib.request.Request(url, headers={
        "Authorization": f"Bearer {token}",
        "Accept": "application/json",
    })
    with urllib.request.urlopen(req) as resp:
        data = json.loads(resp.read())

    fields = data["fields"]
    summary = fields.get("summary", "")
    desc_obj = fields.get("description", {}) or {}
    desc_text = extract_adf_text(desc_obj)
    full_text = f"{summary}\n{desc_text}"

    return {
        "source": "jira",
        "jira_key": ticket_key,
        "title": summary,
        "jira_url": f"{base_url}/browse/{ticket_key}",
        "priority": fields.get("priority", {}).get("name"),
        "status": fields.get("status", {}).get("name"),
        "labels": fields.get("labels", []),
        "environment": detect_environment(full_text),
        "resource_types": detect_resource_types(full_text),
        "raw_text": full_text[:2000],
    }


def extract_adf_text(node: dict | list | str) -> str:
    """Recursively extract plain text from Atlassian Document Format."""
    if isinstance(node, str):
        return node
    if isinstance(node, list):
        return " ".join(extract_adf_text(n) for n in node)
    if isinstance(node, dict):
        if node.get("type") == "text":
            return node.get("text", "")
        content = node.get("content", [])
        return " ".join(extract_adf_text(c) for c in content)
    return ""


def parse_tf_directory(tf_dir: str) -> dict:
    """Scan existing Terraform files to understand current state."""
    tf_path = Path(tf_dir).expanduser()
    if not tf_path.exists():
        print(f"Directory not found: {tf_dir}", file=sys.stderr)
        sys.exit(1)

    resources = {}
    for tf_file in tf_path.rglob("*.tf"):
        text = tf_file.read_text()
        # Find all resource "type" "name" blocks
        for m in re.finditer(r'resource\s+"([^"]+)"\s+"([^"]+)"', text):
            resource_type, resource_name = m.group(1), m.group(2)
            resources.setdefault(resource_type, []).append({
                "name": resource_name,
                "file": str(tf_file.relative_to(tf_path)),
            })

    return {
        "source": "terraform_directory",
        "source_path": str(tf_dir),
        "environment": detect_environment(str(tf_dir)),
        "existing_resources": resources,
        "resource_types": list(resources.keys()),
        "resource_count": sum(len(v) for v in resources.values()),
    }


def main():
    parser = argparse.ArgumentParser(description="Parse infrastructure requirements")
    group = parser.add_mutually_exclusive_group(required=True)
    group.add_argument("--jira", metavar="TICKET_KEY", help="Jira ticket key, e.g. INFRA-123")
    group.add_argument("--md", metavar="PATH", help="Markdown file path")
    group.add_argument("--docx", metavar="PATH", help="DOCX file path")
    group.add_argument("--tf-dir", metavar="PATH", help="Existing Terraform directory to scan")
    parser.add_argument("--pretty", action="store_true", help="Pretty-print output")

    args = parser.parse_args()

    if args.jira:
        result = parse_jira(args.jira)
    elif args.md:
        result = parse_markdown(args.md)
    elif args.docx:
        result = parse_docx(args.docx)
    elif args.tf_dir:
        result = parse_tf_directory(args.tf_dir)

    indent = 2 if args.pretty else None
    print(json.dumps(result, indent=indent))


if __name__ == "__main__":
    main()
