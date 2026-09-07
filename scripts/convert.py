#!/usr/bin/env python3
"""
convert.py — Converts .md files to .docx for reviewer comments.

Usage:
  python convert.py               # convert all discovered .md files
  python convert.py architecture  # convert only matching files (substring)

Output goes to the docx/ folder (created automatically if missing).
Requires pandoc: brew install pandoc
Optional: pip install python-docx  (enables bold headers + font scaling in tables)
"""

import re
import subprocess
import sys
from pathlib import Path

# DOCS = Path(__file__).parent
# print(DOCS)
DOCS = Path.cwd()
# print(DOCS)
OUT  = DOCS / "docx"

SKIP = {"DEPLOYMENT.md"}

HEADER_PT = 12  # table header row font size
BODY_PT   = 10  # table body row font size

# Characters forbidden in XML 1.0 (breaks DOCX readers like Google Docs).
# Allowed: #x9 | #xA | #xD | [#x20-#xD7FF] | [#xE000-#xFFFD] | [#x10000-#x10FFFF]
_FORBIDDEN_XML = re.compile(
    r"[\x00-\x08\x0b\x0c\x0e-\x1f\ud800-\udfff￾￿]"
)


def sanitize(text: str) -> str:
    """Replace XML-forbidden codepoints with their U+XXXX representation."""
    return _FORBIDDEN_XML.sub(lambda m: f"U+{ord(m.group()):04X}", text)


def discover() -> list[Path]:
    # for p in DOCS.glob("*.md"):
    #     print(p.name)
    return sorted(p for p in DOCS.glob("*.md") if p.name not in SKIP)


def check_pandoc() -> None:
    if subprocess.run(["which", "pandoc"], capture_output=True).returncode != 0:
        sys.exit("pandoc not found — install it with: brew install pandoc")


def _set_cell_bg(cell, hex_color: str) -> None:
    from docx.oxml import parse_xml
    from docx.oxml.ns import nsdecls, qn
    tcPr = cell._tc.get_or_add_tcPr()
    for tag in (qn("w:cnfStyle"), qn("w:shd")):
        for el in tcPr.findall(tag):
            tcPr.remove(el)
    fill = hex_color.lstrip("#").upper()
    tcPr.append(parse_xml(
        f'<w:shd {nsdecls("w")} w:val="clear" w:color="auto" w:fill="{fill}"/>'
    ))


def style_tables(path: Path) -> None:
    """Bold header row with background color and scale font sizes in all DOCX tables."""
    try:
        from docx import Document
        from docx.shared import Pt
        from docx.oxml.ns import qn
    except ImportError:
        return

    doc = Document(str(path))
    for table in doc.tables:
        if not table.rows:
            continue
        # Turn off first-row conditional formatting so cell-level shading wins
        tblPr = table._tbl.find(qn("w:tblPr"))
        if tblPr is not None:
            tblLook = tblPr.find(qn("w:tblLook"))
            if tblLook is not None:
                tblLook.set(qn("w:firstRow"), "0")
        for cell in table.rows[0].cells:
            _set_cell_bg(cell, "D9EAD3")
            for para in cell.paragraphs:
                for run in para.runs:
                    run.bold = True
                    run.font.size = Pt(HEADER_PT)
        for row in table.rows[1:]:
            for cell in row.cells:
                for para in cell.paragraphs:
                    for run in para.runs:
                        run.font.size = Pt(BODY_PT)
    doc.save(str(path))


def convert(md: Path) -> None:
    out = OUT / (md.stem + ".docx")
    clean = sanitize(md.read_text(encoding="utf-8"))
    result = subprocess.run(
        ["pandoc", "-f", "markdown", "-o", str(out), "--resource-path", str(DOCS)],
        input=clean, capture_output=True, text=True,
    )
    if result.returncode == 0:
        style_tables(out)
        size = out.stat().st_size
        print(f"  OK    {md.name} → docx/{out.name}  ({size // 1024} KB)")
    else:
        print(f"  FAIL  {md.name}\n        {result.stderr.strip()}")


def main() -> None:
    targets = sys.argv[1:]
    all_files = discover()

    selected = (
        [f for f in all_files if any(t in f.name for t in targets)]
        if targets else all_files
    )

    if not selected:
        print(f"No files matched: {targets}")
        sys.exit(1)

    check_pandoc()
    OUT.mkdir(exist_ok=True)

    print(f"Discovered {len(all_files)} file(s), converting {len(selected)}...")
    for md in selected:
        convert(md)
    print("Done.")


if __name__ == "__main__":
    main()
